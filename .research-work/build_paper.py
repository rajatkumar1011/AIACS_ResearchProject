from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Rajat Kumar\Desktop\KJSSE\SEM-V\HO-AIACS Lab\Research Paper")
SOURCE = ROOT / "IJSRP-paper-format.docx"
OUTPUT = ROOT / "Explainable-AI-IDS-Research-Paper-Revised-Final.docx"
PERFORMANCE_FIG = ROOT / ".research-work" / "experiment-output" / "performance.png"
GLOBAL_SHAP_FIG = ROOT / ".research-work" / "experiment-output" / "global_shap.png"
LOCAL_SHAP_FIG = ROOT / ".research-work" / "experiment-output" / "local_shap_waterfall.png"
LOCAL_LIME_FIG = ROOT / ".research-work" / "experiment-output" / "local_lime_explanation.png"


def set_font(run, name="Times New Roman", size=9.5, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def clear_container(container):
    for child in list(container._element):
        container._element.remove(child)


def remove_all_body_content(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def set_columns(section, count=2, space_twips=288):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(space_twips))


def set_cell_margins(cell, top=45, start=55, bottom=45, end=55):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color="BFBFBF", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True


def add_major_heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    keep_with_next(p)
    set_font(p.add_run(text.upper()), size=10.5)
    return p


def add_subheading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    keep_with_next(p)
    set_font(p.add_run(text), size=10.5, bold=True)
    return p


def add_body(doc, text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(1.6)
    p.paragraph_format.line_spacing = 1.08
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.18)
    set_font(p.add_run(text), size=10.5)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.17)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1)
    set_font(p.add_run("• "), size=10.5)
    set_font(p.add_run(text), size=10.5)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    keep_with_next(p)
    set_font(p.add_run(text), size=8, italic=True)
    return p


def add_figure_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_together = True
    set_font(p.add_run(text), size=8, italic=True)
    return p


def add_table(doc, headers, rows, widths=None, font_size=7.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    hdr_pr = hdr._tr.get_or_add_trPr()
    hdr_pr.append(OxmlElement("w:tblHeader"))
    hdr_pr.append(OxmlElement("w:cantSplit"))
    for i, label in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, "D9E2F3")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(label), size=font_size, bold=True)
        if widths:
            cell.width = Inches(widths[i])
    for row_idx, values in enumerate(rows):
        new_row = table.add_row()
        new_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cells = new_row.cells
        for i, value in enumerate(values):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx % 2:
                set_cell_shading(cell, "F7F9FC")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(str(value)), size=font_size)
            if widths:
                cell.width = Inches(widths[i])
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    return table


def add_reference(doc, number, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(1.3)
    set_font(p.add_run(f"[{number}] "), size=8)
    set_font(p.add_run(text), size=8)
    return p


shutil.copy2(SOURCE, OUTPUT)
doc = Document(OUTPUT)
remove_all_body_content(doc)

# Core metadata
props = doc.core_properties
props.title = "Explainable AI for Cybersecurity Interpreting Machine Learning Decisions in Intrusion Detection Systems"
props.subject = "AI Applications in Cyber Security research paper"
props.author = "Rajat Kumar; Sadhil Madan"
props.keywords = "explainable artificial intelligence, intrusion detection, SHAP, LIME, UNSW-NB15"
props.comments = "Completed course research paper with reproducible UNSW-NB15 experiments and explainability analysis."

# Normalize styles used in the rebuilt body.
if "Title" not in doc.styles:
    doc.styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
for style_name in ("Normal", "Title"):
    if style_name in doc.styles:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
        style.font.color.rgb = RGBColor(0, 0, 0)

first = doc.sections[0]
first.page_width = Inches(8.5)
first.page_height = Inches(11)
first.left_margin = Inches(0.5)
first.right_margin = Inches(0.5)
first.top_margin = Inches(0.55)
first.bottom_margin = Inches(0.55)
first.header_distance = Inches(0.25)
first.footer_distance = Inches(0.25)
set_columns(first, 1)

# Header and footer
clear_container(first.header)
hp = first.header.add_paragraph()
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hp.paragraph_format.space_after = Pt(0)
set_font(hp.add_run("Explainable AI for Cybersecurity"), size=8)

clear_container(first.footer)
fp = first.footer.add_paragraph()
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.space_after = Pt(0)
set_font(fp.add_run("AI Applications in Cyber Security   |   Page "), size=8)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
r = OxmlElement("w:r")
rpr = OxmlElement("w:rPr")
rfonts = OxmlElement("w:rFonts")
rfonts.set(qn("w:ascii"), "Times New Roman")
rfonts.set(qn("w:hAnsi"), "Times New Roman")
sz = OxmlElement("w:sz")
sz.set(qn("w:val"), "16")
rpr.append(rfonts)
rpr.append(sz)
r.append(rpr)
t = OxmlElement("w:t")
t.text = "1"
r.append(t)
fld.append(r)
fp._p.append(fld)

# Title block
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(9)
p.paragraph_format.space_after = Pt(7)
set_font(p.add_run("Explainable AI for Cybersecurity: Interpreting Machine Learning Decisions in Intrusion Detection Systems"), size=18, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(3)
set_font(p.add_run("Rajat Kumar* and Sadhil Madan**"), size=11, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(5)
set_font(p.add_run("Department of Information Technology, K. J. Somaiya School of Engineering, Mumbai, India\n"), size=9)
set_font(p.add_run("*ORCID: 0009-0004-1688-0855"), size=8.5)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(3)
r = set_font(p.add_run("Abstract—"), size=9, bold=True, italic=True)
set_font(p.add_run(
    "Machine-learning intrusion detection systems can classify large volumes of network traffic, but their predictions often provide little evidence that a security analyst can examine. This study compares Random Forest (RF) and XGBoost (XGB) on the official UNSW-NB15 training and test partitions for binary attack detection, then evaluates SHAP and LIME explanations across true positives, true negatives, false positives, and false negatives. At a fixed 0.5 threshold, Random Forest obtained F1=0.9055, recall=0.9818, FPR=0.2287, and ROC-AUC=0.9840; XGBoost obtained F1=0.8962, recall=0.9833, FPR=0.2587, and ROC-AUC=0.9840. Random Forest therefore provided the stronger precision and false-alarm trade-off, while XGBoost missed 65 fewer attacks. Global SHAP rankings shared five of their ten leading attributes but had only moderate overall rank correlation. Median top-five SHAP-LIME overlap ranged from 0.111 to 0.429, showing that local explanations often emphasized different evidence. TreeSHAP was substantially faster than LIME, while LIME fidelity was higher for XGBoost. These results support explanation as diagnostic evidence rather than proof that an alert is correct or causal."
), size=9)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.space_after = Pt(4)
set_font(p.add_run("Index Terms—"), size=9, bold=True, italic=True)
set_font(p.add_run("cybersecurity, explainable artificial intelligence, intrusion detection system, LIME, SHAP"), size=9)

# Begin two-column body.
body_section = doc.add_section(WD_SECTION.CONTINUOUS)
body_section.top_margin = Inches(0.62)
body_section.bottom_margin = Inches(0.55)
body_section.left_margin = Inches(0.5)
body_section.right_margin = Inches(0.5)
body_section.header_distance = Inches(0.25)
body_section.footer_distance = Inches(0.25)
set_columns(body_section, 2)

add_major_heading(doc, "I. Introduction")
add_body(doc, "Intrusion detection systems examine network events for signs of misuse, attack, or policy violation. Machine-learning models extend this task beyond fixed signatures by learning statistical patterns from labelled traffic. This flexibility is valuable when attack behaviour varies, but it also creates an operational problem: a model may label a flow as malicious without giving the analyst a reason that can be checked against the traffic context.")
add_body(doc, "An unexplained false positive can waste investigation time or encourage alert fatigue. An unexplained false negative is more serious because harmful activity may pass without review. Accuracy alone hides this difference. A useful evaluation must therefore report recall and false-positive rate, preserve the natural test distribution, and examine why correct and incorrect predictions receive their scores.")
add_body(doc, "Explainable artificial intelligence provides methods for describing a trained model's behaviour. SHAP assigns contributions to input features through an additive framework derived from Shapley values [3]. LIME fits a simple surrogate around one prediction after creating nearby samples [2]. Both techniques can help expose influential attributes, but neither proves that a prediction is correct or that an attributed feature caused an attack. Explanations can also change with reference data, perturbation rules, or random seeds.")
add_body(doc, "This paper proposes a controlled comparison of Random Forest and XGBoost on UNSW-NB15. It evaluates detection and explanation under one data protocol, then separates explanation analysis by true positive, true negative, false positive, and false negative outcomes. The central question is: How effectively can Explainable AI techniques improve the interpretability of machine-learning-based intrusion detection systems while maintaining reliable intrusion-detection performance?")
add_body(doc, "The study uses interpretability in a limited, testable sense. An explanation is more useful for this experiment when it identifies a compact set of understandable network attributes, approximates the model faithfully near the selected record, remains reasonably stable under controlled reruns, and can be generated within a recorded time. These properties do not by themselves show that an analyst understands the alert or will make a better response decision. That broader claim requires a separate human evaluation.")
add_body(doc, "The paper contributes a reproducible full-split benchmark, an error-stratified SHAP-LIME comparison, and a practical account of explanation agreement, fidelity, stability, and time. The implementation files retain the trained pipelines, generated figures, and machine-readable measurements so that every reported number can be checked against the same run.")

add_major_heading(doc, "II. Background and Related Work")
add_subheading(doc, "A. Explainability Methods")
add_body(doc, "Explainability can describe an entire model or one decision. Global explanations summarise behaviour across many records, for example by ranking features using mean absolute attribution. Local explanations describe why a particular record received a score. IDS evaluation needs both levels: global summaries can reveal persistent reliance on a feature, while local explanations can show why a single alert differs from the usual pattern. A global ranking should not be used as though it were the explanation for every alert.")
add_body(doc, "LIME explains a prediction by sampling around the selected record, obtaining model outputs for those samples, and fitting an interpretable local model [2]. Its model-agnostic design is useful when the same explanation procedure must be applied to different classifiers. The local approximation can nevertheless become unstable when sampling changes or when perturbations create implausible combinations of network attributes.")
add_body(doc, "The choices made around LIME are part of the experiment rather than minor software settings. Kernel width controls how strongly distant samples influence the surrogate. The number of perturbed samples affects runtime and the precision of the fitted local model. Because UNSW-NB15 includes related counts, durations, rates, and protocol fields, perturbations may be mathematically valid but unlike feasible flows. The implementation preserved categorical values and interpreted LIME as a local approximation.")
add_body(doc, "SHAP represents a prediction as a baseline plus feature contributions [3]. Tree-specific implementations can efficiently explain tree ensembles such as Random Forest and XGBoost. A SHAP summary can describe global model reliance, while waterfall or bar plots can explain one alert. Attributions remain conditional on the model, feature representation, background sample, and dependence assumptions. Correlated traffic features can divide or redistribute importance in ways that require cautious interpretation.")
add_body(doc, "For a prediction f(x), an additive explanation can be written conceptually as f(x) = φ0 + Σφi, where φ0 is the reference output and φi is the contribution assigned to feature i. Positive and negative contributions are meaningful only relative to the selected class and output scale. This study used raw-output TreeSHAP values and probability-based LIME surrogates, so attribution magnitudes were not compared directly across methods.")

add_subheading(doc, "B. Explainable Intrusion Detection")
add_body(doc, "Wang et al. presented an early explainable machine-learning framework for IDS using SHAP with one-versus-all and multiclass classifiers on NSL-KDD [4]. The explanations corresponded to known attack characteristics, establishing a practical link between feature attribution and intrusion analysis. The use of an older benchmark and the limited treatment of explanation reliability leave room for evaluation on other traffic representations.")
add_body(doc, "Patil et al. trained decision tree, Random Forest, SVM, and voting models on CICIDS2017 and used LIME to explain individual results [5]. Their study shows that ensemble detection and local explanation can be integrated. It concentrates on LIME and reported classifier performance, so agreement with another explainer and a systematic comparison of false positives and false negatives remain separate questions.")
add_body(doc, "Wei et al. developed xNIDS to explain deep network intrusion models while accounting for input history and feature dependencies [6]. The framework evaluates fidelity, sparsity, completeness, and stability and can convert explanations into defence rules. This work demonstrates that explanation quality has measurable dimensions beyond visual appeal, although its deep-model setting differs from the tabular tree ensembles considered here.")
add_body(doc, "Arreche et al. proposed XAI-IDS and evaluated seven classifiers with SHAP and LIME on CICIDS2017, NSL-KDD, and RoEduNet-SIMARGL2021 [7]. Their framework compares model behaviour, feature selection, and explanation techniques across datasets. It confirms that broad multi-model XAI comparisons already exist; a defensible contribution for the present project must therefore come from its controlled UNSW-NB15 protocol and error-stratified analysis.")
add_body(doc, "Gaspar et al. applied LIME and SHAP to an MLP-based IoT intrusion solution using a generated form of the ADFA-LD system-call dataset [8]. They used perturbation analysis to test explanations and collected participant feedback. Their work supports explicit evaluation of explanations, while its system-call features and neural model limit direct transfer to network-flow tree models.")
add_body(doc, "Mia et al. compared SHAP patterns for correct predictions and errors across CIC-IoT2023, NF-UQ-NIDS-v2, and HIKARI-2021 [9]. Their overlapping plots are designed to guide analysts in identifying potential false positives and false negatives. The approach gives the current study a direct methodological precedent; the planned extension compares a second explainer, retains the original test distribution, and reports explanation time and stability.")
add_body(doc, "Hermosilla et al. compared SHAP and LIME for XGBoost and TabNet on UNSW-NB15, including stability and fidelity considerations in a forensic setting [10]. This is the closest prior configuration and prevents a broad novelty claim based only on the dataset and explainers. The present study instead pairs XGBoost with Random Forest and organises the comparison around all four confusion-matrix outcomes.")
add_body(doc, "Mohale and Obagbuwa's systematic review found continued problems in standardisation, scalability, and real-time use of explanations in IDS [11]. Taken together, the literature shows that explanations should be evaluated as outputs with their own limitations. It also shows that explanation plots alone do not establish human understanding or operational trust.")
add_body(doc, "The studies also use different meanings of explanation quality. Some judge whether a plot appears consistent with known attack characteristics. Others perturb important features, compare a surrogate with the original model, test stability, or ask participants to rate usefulness. These measures are complementary rather than interchangeable. The present study therefore uses several measurable properties and labels each one precisely. It does not collapse fidelity, stability, agreement, sparsity, runtime, and human trust into one interpretability score.")
add_body(doc, "Table I summarises the relationship between these studies and the controlled replication and extension reported here.")

# Full-width literature table.
full = doc.add_section(WD_SECTION.CONTINUOUS)
full.top_margin = Inches(0.62)
full.bottom_margin = Inches(0.55)
full.left_margin = Inches(0.5)
full.right_margin = Inches(0.5)
set_columns(full, 1)
add_caption(doc, "Table I. Selected literature and its relevance to the present study")
add_table(doc,
    ["Study", "Method", "Dataset", "Main finding", "Limitation or gap"],
    [
        ["Wang et al. [4]", "SHAP with neural IDS", "NSL-KDD", "Local and global attributions align with attack characteristics", "Older dataset; limited explanation validation"],
        ["Patil et al. [5]", "Ensemble IDS with LIME", "CICIDS2017", "Local explanations accompany an ensemble classifier", "Single explainer; limited error-stratified analysis"],
        ["Wei et al. [6]", "xNIDS", "Four DL-NIDS settings", "Measures fidelity, sparsity, completeness, and stability", "Deep sequential setting differs from tabular tree models"],
        ["Arreche et al. [7]", "Seven models; SHAP and LIME", "CICIDS2017, NSL-KDD, RoEduNet", "Broad model and explainer comparison", "Does not centre on UNSW-NB15 error outcomes"],
        ["Gaspar et al. [8]", "MLP; LIME and SHAP", "Generated ADFA-LD", "Perturbation and user feedback evaluate explanations", "System-call data; one model family"],
        ["Mia et al. [9]", "SHAP error-group plots", "Three traffic datasets", "Diagnoses FP and FN patterns", "No LIME comparison; visual similarity depends on analyst judgement"],
        ["Hermosilla et al. [10]", "XGBoost, TabNet; SHAP and LIME", "UNSW-NB15", "Compares fidelity and stability for forensic use", "Different model pairing; broader outcome-stratified replication is useful"],
        ["Mohale and Obagbuwa [11]", "Systematic review", "Twenty studies", "Finds standardisation and scalability gaps", "Review evidence does not replace controlled experiments"],
    ],
    widths=[0.95, 1.45, 1.22, 2.2, 2.15], font_size=6.8)

back = doc.add_section(WD_SECTION.CONTINUOUS)
back.top_margin = Inches(0.62)
back.bottom_margin = Inches(0.55)
back.left_margin = Inches(0.5)
back.right_margin = Inches(0.5)
set_columns(back, 2)

add_major_heading(doc, "III. Problem Statement and Research Gap")
add_body(doc, "High-performing IDS classifiers can still be difficult to audit. A model comparison based only on accuracy cannot show whether one system misses more attacks, creates more false alarms, or relies on features that are unstable under a different explanation sample. Existing research has already applied SHAP and LIME to IDS, compared multiple models, and studied misclassifications [7]–[10]. The research gap is therefore specific rather than absolute.")
add_body(doc, "The contribution is a reproducible comparison of Random Forest and XGBoost under one UNSW-NB15 protocol. Detection metrics and explanation measures are reported together. Explanations are compared across the same records where possible and stratified by confusion-matrix outcome. The study also records agreement, stability, fidelity, and processing time. This design treats the work as a scoped replication and extension, not a new XAI algorithm.")
add_body(doc, "Three controls make the comparison useful. First, both classifiers receive the same predictor definitions, training records, validation logic, and final test records. Second, the original test distribution is preserved so that false-positive rate and precision describe the benchmark rather than a rebalanced sample. Third, explanations are aggregated from transformed columns to original network attributes before rankings are compared. Without these controls, differences may reflect preprocessing or presentation choices instead of the models.")
add_body(doc, "The study's practical question is diagnostic. If a model flags benign traffic, do its explanations resemble typical attack detections or expose an unusual reliance on a small number of features? If a model misses an attack, which attributes pushed the score toward benign traffic, and does the competing model use those attributes differently? These questions do not promise automatic error correction. They define evidence that can guide later review of data quality, features, thresholds, and model design.")

add_major_heading(doc, "IV. Research Questions and Objectives")
add_subheading(doc, "A. Research Questions")
add_body(doc, "RQ1: How do Random Forest and XGBoost compare in attack recall, false-positive rate, F1-score, ROC-AUC, and PR-AUC under identical data partitions?", indent=False)
add_body(doc, "RQ2: Which original network attributes most strongly influence each model, and how much do their global feature rankings overlap?", indent=False)
add_body(doc, "RQ3: How do SHAP and LIME explanations differ among true positives, true negatives, false positives, and false negatives?", indent=False)
add_body(doc, "RQ4: How repeatable are the explanations under controlled reruns, and what computation time do they require?", indent=False)
add_body(doc, "RQ5: Which diagnostic observations can be made from error explanations, and which claims about analyst trust remain untested without a user study?", indent=False)

add_subheading(doc, "B. Objectives")
add_bullet(doc, "Build a leakage-controlled preprocessing and training pipeline using the official UNSW-NB15 partitions.")
add_bullet(doc, "Compare Random Forest and XGBoost using the same test records and clearly defined positive class.")
add_bullet(doc, "Generate global and local explanations with network attributes presented in their original semantic groups.")
add_bullet(doc, "Analyse correct predictions, false alarms, and missed attacks, then measure explanation agreement, sensitivity, fidelity, and time.")

add_major_heading(doc, "V. Methodology and Implementation")
add_body(doc, "The implemented sequence comprised dataset acquisition, schema audit, preprocessing, training-only model selection, final model fitting, intrusion prediction, performance evaluation, SHAP and LIME explanation, and error-group comparison. The scripts save fitted pipelines, machine-readable result files, figures, case identifiers, random seeds, and software versions.")
add_body(doc, "The predictor matrix contained 42 attributes. The identifier, binary label, and attack-category label were excluded; the attack category was retained only for error analysis. Protocol, service, and state were one-hot encoded with unknown-category handling, while numeric features were passed to the tree models without scaling. No test labels informed preprocessing or fitting.")
add_body(doc, "Figure 1 summarises the implemented workflow from the dataset audit to the comparison of detection and explanation results.")

add_caption(doc, "Fig. 1. Implemented experimental and explanation workflow")
add_table(doc, ["Stage", "Output"], [
    ["1. Audit and split", "Verified schema, counts, labels, and untouched test partition"],
    ["2. Preprocess", "Training-fitted cleaning, encoding, and feature map"],
    ["3. Train", "Tuned Random Forest and XGBoost pipelines"],
    ["4. Evaluate", "Probabilities, confusion outcomes, and detection metrics"],
    ["5. Explain", "SHAP and LIME outputs for global and outcome-group samples"],
    ["6. Compare", "Agreement, stability, fidelity, runtime, and error interpretation"],
], widths=[1.0, 2.6], font_size=7.2)

add_subheading(doc, "A. Dataset Description")
add_body(doc, "UNSW-NB15 was generated in the Cyber Range Lab at UNSW Canberra using IXIA PerfectStorm to combine normal activity with synthetic contemporary attacks [1]. The complete collection contains 2,540,044 records and 49 attributes including the class label. Its published training and test files contain 175,341 and 82,332 records, respectively. Attack categories include Analysis, Backdoors, DoS, Exploits, Fuzzers, Generic, Reconnaissance, Shellcode, and Worms [1].")
add_body(doc, "The primary task is binary classification, with attack as the positive class. The attack category is excluded from predictors and retained only for error reporting. The fixed test partition supports repeatability, while its laboratory origin and age restrict claims about modern production traffic.")
add_body(doc, "The feature set describes a flow from several perspectives. Basic attributes cover protocol, service, state, duration, bytes, and packet counts. Content and time-related attributes describe connection behaviour, while additional generated features summarise related activity involving hosts, services, ports, and timing [1]. These groups create useful semantic material for explanation, but they also contain dependencies. Duration, rate, bytes, and packet counts cannot be treated as independent physical quantities when local samples are generated.")
add_body(doc, "Class imbalance affects both model fitting and interpretation. The binary task contains many attacks overall while some categories remain rare. A high aggregate recall can therefore coexist with weaker recall in a category such as Fuzzers. Category errors are reported as counts and rates, with restrained interpretation for small denominators.")
add_body(doc, "Table II records the dataset partitions and the role assigned to each field group in the experiment.")
add_caption(doc, "Table II. UNSW-NB15 study profile")
add_table(doc, ["Property", "Implemented use"], [
    ["Training records", "175,341; preprocessing, selection, and fitting"],
    ["Test records", "82,332; final evaluation only"],
    ["Prediction task", "Benign versus attack"],
    ["Attack category", "Evaluation metadata for missed-attack analysis"],
    ["Feature types", "Numeric and categorical network-flow attributes"],
], widths=[1.05, 2.5], font_size=7.4)

add_subheading(doc, "B. Data Preprocessing and Feature Selection")
add_body(doc, "The supplied files contained 175,341 training rows and 82,332 test rows with the expected 45 columns. The training labels comprised 119,341 attacks and 56,000 benign flows; the test labels comprised 45,332 attacks and 37,000 benign flows. The identifier, label, and attack category did not enter either predictor.")
add_body(doc, "Protocol, service, and state were one-hot encoded by a transformer fitted only on training data; unseen test categories were ignored rather than causing failure. The remaining 39 numeric features were passed through unchanged because the selected tree models do not require standardisation. The same fitted transformer was embedded in both saved model pipelines.")
add_body(doc, "The experiment retained all 42 predictors and the original class distribution. No resampling, test-set balancing, or test-driven feature selection was performed. This full-feature design keeps the comparison focused on the classifiers and explainers and prevents a feature-selection loop from contaminating the final evaluation.")

add_subheading(doc, "C. Machine Learning Models")
add_body(doc, "Random Forest builds an ensemble of decision trees from bootstrapped samples and random feature subsets [12]. Averaging across trees can reduce the variance of a single tree, while class weights and depth controls help manage imbalance and complexity. XGBoost adds trees sequentially to correct current errors and uses regularisation, shrinkage, and sampling to control overfitting [13]. Both models support probability estimates and tree-based SHAP, which permits a consistent explanation interface.")
add_body(doc, "A fixed 20% stratified validation subset of the training file compared three configurations per model using seed 42. The selected Random Forest used 220 trees, unrestricted depth, minimum leaf size two, square-root feature sampling, and balanced bootstrap weights. The selected XGBoost model used 260 trees, depth six, learning rate 0.08, and row and column subsampling of 0.85.")
add_body(doc, "Both final pipelines were refitted on all training records and evaluated at the declared 0.5 decision threshold. Continuous scores were retained for ROC-AUC and PR-AUC. No probability calibration was added, so probability magnitudes are interpreted as model scores rather than calibrated event frequencies.")

add_subheading(doc, "D. Explainable AI Techniques")
add_body(doc, "Approximate TreeSHAP produced raw-output attributions for each tree ensemble. Contributions from one-hot columns were summed back to their original attributes. Global rankings used a label-stratified sample of 1,000 test records, containing 500 benign and 500 attack flows.")
add_body(doc, "LIME called each complete preprocessing-and-model pipeline. Categorical perturbations were constrained to observed protocol, service, and state values. Each explanation used 1,500 perturbations. Its local surrogate R² was retained as fidelity evidence rather than assuming the surrogate matched the classifier.")
add_body(doc, "Local comparison used 15 seeded records from each of TP, TN, FP, and FN for each model. The top-five feature-set Jaccard score measured SHAP-LIME overlap. LIME stability used three seeds on five records per outcome. TreeSHAP repeatability was checked under identical model and input settings; because the implemented tree-path procedure is deterministic, this is a repeatability check rather than a background-sensitivity experiment.")

add_subheading(doc, "E. Experimental Setup")
add_body(doc, "The implementation used Python 3.12.14, scikit-learn 1.9.1, XGBoost 3.4.1, SHAP 0.52.0, LIME 0.2.0.1, and seed 42 on Windows 11. The official test file remained outside model selection. The 0.5 threshold was fixed before test scoring.")
add_body(doc, "All reported samples were selected by code from stored predictions. The balanced explanation sample supports comparisons among TP, TN, FP, and FN but does not estimate their frequency in natural traffic. Runtime was measured for the implemented batch SHAP call and individual LIME calls on the same machine; the resulting values are comparative rather than hardware-independent benchmarks.")

add_subheading(doc, "F. Evaluation Metrics")
add_body(doc, "Accuracy, precision, recall, F1, false-positive rate, ROC-AUC, and PR-AUC measure detection. For binary labels, precision = TP/(TP+FP), recall = TP/(TP+FN), F1 is their harmonic mean, and false-positive rate = FP/(FP+TN). Confusion counts and attack-category misses provide the denominators hidden by aggregate rates.")
add_body(doc, "Explanation comparison used top-five Jaccard overlap, directional agreement on shared features, LIME surrogate R², seeded LIME stability, deterministic SHAP repeatability, and median time per record. These measures answer different questions; agreement or stability does not prove that an explanation is true.")

# Full-width results block for tables and figures.
results_full = doc.add_section(WD_SECTION.CONTINUOUS)
results_full.top_margin = Inches(0.62)
results_full.bottom_margin = Inches(0.55)
results_full.left_margin = Inches(0.5)
results_full.right_margin = Inches(0.5)
set_columns(results_full, 1)

add_major_heading(doc, "VI. Results and Analysis")
add_caption(doc, "Table III. Test performance on 82,332 records at threshold 0.5")
add_table(doc, ["Model", "Accuracy", "Prec.", "Recall", "F1", "FPR", "ROC-AUC", "PR-AUC"], [
    ["Random Forest", "0.8872", "0.8403", "0.9818", "0.9055", "0.2287", "0.9840", "0.9880"],
    ["XGBoost", "0.8745", "0.8232", "0.9833", "0.8962", "0.2587", "0.9840", "0.9884"],
], widths=[1.1, 0.78, 0.68, 0.68, 0.62, 0.62, 0.76, 0.76], font_size=7.2)
add_body(doc, "Table III reports the fixed-threshold metrics. Random Forest produced TN=28,538, FP=8,462, FN=823, and TP=44,509. XGBoost produced TN=27,428, FP=9,572, FN=758, and TP=44,574. XGBoost detected 65 additional attacks, but Random Forest avoided 1,110 false alarms and therefore achieved higher precision, F1, and accuracy. Their ROC-AUC values were nearly identical; PR-AUC slightly favoured XGBoost. Figure 2 shows the corresponding ROC, precision-recall, and fixed-threshold comparisons.")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(1)
p.add_run().add_picture(str(PERFORMANCE_FIG), width=Inches(7.1))
add_figure_caption(doc, "Fig. 2. ROC, precision-recall, and fixed-threshold performance on the official test split")

add_subheading(doc, "A. Explainability Analysis")
add_body(doc, "Figure 3 presents the global SHAP analysis. In the 1,000-record sample, sttl ranked first for both models. Random Forest next emphasized ct_state_ttl, sload, state, and ct_srv_dst; XGBoost emphasized proto, service, dbytes, and response_body_len. Five attributes appeared in both top tens: sttl, service, dbytes, dttl, and ct_srv_dst. The top-ten Jaccard score was 0.333 and the full-ranking Spearman correlation was 0.483, indicating related but different reliance patterns.")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(1)
p.add_run().add_picture(str(GLOBAL_SHAP_FIG), width=Inches(7.1))
add_figure_caption(doc, "Fig. 3. Global mean absolute approximate TreeSHAP values on 1,000 stratified test records")

add_caption(doc, "Table IV. Outcome-stratified local explanation measurements (median values)")
add_table(doc, ["Model", "Outcome", "n", "Top-5 Jacc.", "LIME R²", "LIME stab.", "SHAP ms", "LIME ms"], [
    ["RF", "TP", "15", "0.250", "0.637", "0.667", "3.99", "105.44"],
    ["RF", "TN", "15", "0.429", "0.632", "0.667", "3.99", "98.13"],
    ["RF", "FP", "15", "0.111", "0.699", "0.667", "3.99", "99.78"],
    ["RF", "FN", "15", "0.111", "0.722", "0.667", "3.99", "101.76"],
    ["XGB", "TP", "15", "0.250", "0.801", "0.429", "1.16", "61.24"],
    ["XGB", "TN", "15", "0.250", "0.787", "0.429", "1.16", "65.80"],
    ["XGB", "FP", "15", "0.111", "0.776", "0.429", "1.16", "55.32"],
    ["XGB", "FN", "15", "0.250", "0.779", "0.429", "1.16", "56.63"],
], widths=[0.52, 0.58, 0.3, 0.82, 0.7, 0.76, 0.62, 0.62], font_size=6.8)
add_body(doc, "Table IV shows that top-five agreement was modest and fell to 0.111 for false positives under both models and for Random Forest false negatives. When SHAP and LIME shared a leading feature, median directional agreement was 1.0 for TP and TN, but it fell to 0.5 for FN in both models. XGBoost's LIME surrogates had higher median fidelity (0.776–0.801) than Random Forest's (0.632–0.722). LIME stability was 0.667 for Random Forest and 0.429 for XGBoost. Identical TreeSHAP calls were deterministic, but this does not test sensitivity to a different attribution convention or background distribution.")

add_subheading(doc, "B. Representative Local Explanation")
add_body(doc, "Figure 4 compares real explanations for the same Random Forest true-positive prediction: test row 4,254, with an attack score of 0.9959. The SHAP waterfall highlighted sload, state, rate, ct_dst_sport_ltm, and sttl. LIME instead emphasized proto, sttl, ct_state_ttl, dttl, and state, with a local surrogate R² of 0.620. Only sttl and state appeared in both top-five sets, showing why one explanation method should not be treated as the single correct account of a prediction.")
local_table = doc.add_table(rows=1, cols=2)
local_table.alignment = WD_TABLE_ALIGNMENT.CENTER
local_table.autofit = False
remove_table_borders(local_table)
for cell, image_path in zip(local_table.rows[0].cells, [LOCAL_SHAP_FIG, LOCAL_LIME_FIG]):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=0, start=20, bottom=0, end=20)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(image_path), width=Inches(3.55))
add_figure_caption(doc, "Fig. 4. Local explanations for the same Random Forest true-positive prediction: (a) SHAP waterfall and (b) LIME feature contributions")

add_subheading(doc, "C. False Positive and False Negative Analysis")
add_body(doc, "A false positive is a benign flow incorrectly labelled as an attack. Operationally, it creates an unnecessary alert and consumes analyst time; a sustained high false-positive rate can produce alert fatigue and delay investigation of real attacks. A false negative is an attack incorrectly labelled as benign. It is more dangerous because malicious traffic may pass without review. XAI does not fix either error automatically, but it helps an analyst inspect which observed features influenced the model's decision.")
add_body(doc, "The error distribution exposes the model trade-off. Random Forest misclassified 22.87% of benign test flows as attacks, compared with 25.87% for XGBoost. Most missed attacks were Fuzzers: 759 of 6,062 for Random Forest (12.52%) and 636 of 6,062 for XGBoost (10.49%). Random Forest had 56 missed Exploits and only eight missed attacks across the remaining reported categories; XGBoost had 63 missed Exploits, 21 DoS, 19 Analysis, six Backdoors, seven Generic, four Shellcode, and two Reconnaissance cases. No Worms were missed by either model in this run.")
add_body(doc, "The low SHAP-LIME overlap for error cases means that a single local explanation should not be treated as a settled account of an error. SHAP describes the fitted tree ensemble on its raw-output scale, while LIME approximates the local probability surface created by its perturbations. Their disagreement is therefore evidence about method dependence. Feature attributions remain associations with model output and do not establish that a traffic attribute caused an attack.")

results_back = doc.add_section(WD_SECTION.CONTINUOUS)
results_back.top_margin = Inches(0.62)
results_back.bottom_margin = Inches(0.55)
results_back.left_margin = Inches(0.5)
results_back.right_margin = Inches(0.5)
set_columns(results_back, 2)

add_major_heading(doc, "VII. Discussion")
add_body(doc, "The results operate on two levels. Random Forest and XGBoost achieved recall values of 0.9818 and 0.9833, respectively, but their FPR values remained substantial at 0.2287 and 0.2587. Random Forest gives the better fixed-threshold balance; XGBoost gives the highest recall by a small margin. Explanation metrics show that the account presented to an analyst changes with the explainer, especially for errors.")
add_body(doc, "Post-hoc explanation does not change a fixed model's predictions. Adding SHAP or LIME therefore cannot directly improve accuracy, recall, or false-positive rate. Explanations may support later model redesign, feature review, or threshold adjustment, but each change requires validation on data that did not inform the change. Once test cases guide redesign, those cases are no longer an untouched final evaluation set.")
add_body(doc, "Claims about trust and practical usability require evidence from security analysts. The present benchmark can demonstrate feature visibility, explanation consistency, and processing cost. It cannot demonstrate faster triage, better analyst decisions, or greater trust without a user study. This boundary is important because readable plots may still omit dependencies or reflect artefacts in the data.")
add_body(doc, "Operational preference determines the model choice. Random Forest reduced false alarms by 1,110 relative to XGBoost, which matters when benign volume is high. XGBoost detected 65 additional attacks and missed fewer Fuzzers, which may matter when missed activity carries the dominant cost. Neither fixed threshold is deployment-ready without site-specific calibration.")
add_body(doc, "Approximate TreeSHAP required about 1.16–3.99 ms per sampled record, whereas LIME required about 55–105 ms. This supports batch SHAP for broad review and selective LIME use when a second local perspective is valuable. The low overlap between methods should be displayed as uncertainty, not hidden by choosing the more persuasive explanation.")

add_major_heading(doc, "VIII. Limitations")
add_body(doc, "UNSW-NB15 contains laboratory-generated traffic and may not represent present production networks, encrypted applications, or current attack behaviour. Its predefined split supports repeatability but does not test temporal generalisation or cross-network transfer. Binary labels merge distinct attacks, and rare categories may yield too few errors for stable subgroup analysis.")
add_body(doc, "SHAP and LIME explain model behaviour rather than network causality. Correlated features, one-hot aggregation, approximate TreeSHAP, and LIME sampling can alter attribution. The outcome groups contain only 15 records each for explanation comparison, and the global sample contains 1,000 records. SHAP repeatability was tested under identical calls rather than alternative backgrounds. No confidence intervals or analyst study were performed, so human usefulness remains untested.")

add_major_heading(doc, "IX. Future Scope")
add_body(doc, "After the binary benchmark is complete, the protocol can be extended to multiclass attack recognition, temporal validation, and cross-dataset testing on CICIDS2017 or a newer operational collection. Counterfactual explanations could complement feature attribution by showing feasible changes associated with a different decision. A user study with security analysts could measure decision accuracy, review time, confidence calibration, and the effect of explanation uncertainty. Runtime profiling could also test selective explanation of high-risk alerts rather than every flow.")

add_major_heading(doc, "X. Conclusion")
add_body(doc, "This study completed a reproducible Random Forest and XGBoost comparison on the official UNSW-NB15 partitions. Random Forest achieved the stronger fixed-threshold balance with F1=0.9055 and FPR=0.2287, while XGBoost achieved slightly higher recall at 0.9833. SHAP exposed overlapping global reliance on sttl, service, dbytes, dttl, and ct_srv_dst, yet the models' complete rankings were only moderately correlated. Local SHAP-LIME overlap was low, especially for false positives, and LIME was markedly slower. XAI did not improve the fixed models' accuracy; it made their decisions easier to inspect. Because explanation content depended on the method, these outputs should be used as diagnostic evidence rather than proof of correctness or causality.")

add_major_heading(doc, "References")
refs = [
    "N. Moustafa and J. Slay, “UNSW-NB15: A comprehensive data set for network intrusion detection systems (UNSW-NB15 network data set),” in 2015 Military Communications and Information Systems Conference (MilCIS), 2015, pp. 1–6, doi: 10.1109/MilCIS.2015.7348942.",
    "M. T. Ribeiro, S. Singh, and C. Guestrin, “Why should I trust you?: Explaining the predictions of any classifier,” in Proc. 22nd ACM SIGKDD Int. Conf. Knowledge Discovery and Data Mining, 2016, pp. 1135–1144, doi: 10.1145/2939672.2939778.",
    "S. M. Lundberg and S.-I. Lee, “A unified approach to interpreting model predictions,” in Advances in Neural Information Processing Systems, vol. 30, 2017, pp. 4765–4774.",
    "M. Wang, K. Zheng, Y. Yang, and X. Wang, “An explainable machine learning framework for intrusion detection systems,” IEEE Access, vol. 8, pp. 73127–73141, 2020, doi: 10.1109/ACCESS.2020.2988359.",
    "S. Patil et al., “Explainable artificial intelligence for intrusion detection system,” Electronics, vol. 11, no. 19, Art. no. 3079, 2022, doi: 10.3390/electronics11193079.",
    "F. Wei, H. Li, Z. Zhao, and H. Hu, “xNIDS: Explaining deep learning-based network intrusion detection systems for active intrusion responses,” in 32nd USENIX Security Symposium, 2023, pp. 4337–4354.",
    "O. Arreche, T. Guntur, and M. Abdallah, “XAI-IDS: Toward proposing an explainable artificial intelligence framework for enhancing network intrusion detection systems,” Applied Sciences, vol. 14, no. 10, Art. no. 4170, 2024, doi: 10.3390/app14104170.",
    "D. Gaspar, P. Silva, and C. Silva, “Explainable AI for intrusion detection systems: LIME and SHAP applicability on multi-layer perceptron,” IEEE Access, vol. 12, pp. 30164–30175, 2024, doi: 10.1109/ACCESS.2024.3368377.",
    "M. Mia, M. M. A. Pritom, T. Islam, and K. Hasan, “Visually analyze SHAP plots to diagnose misclassifications in ML-based intrusion detection,” in 2024 IEEE International Conference on Data Mining Workshops, 2024, pp. 632–641, doi: 10.1109/ICDMW65004.2024.00088.",
    "P. Hermosilla, S. Berríos, and H. Allende-Cid, “Explainable AI for forensic analysis: A comparative study of SHAP and LIME in intrusion detection models,” Applied Sciences, vol. 15, no. 13, Art. no. 7329, 2025, doi: 10.3390/app15137329.",
    "V. Z. Mohale and I. C. Obagbuwa, “A systematic review on the integration of explainable artificial intelligence in intrusion detection systems to enhancing transparency and interpretability in cybersecurity,” Frontiers in Artificial Intelligence, vol. 8, Art. no. 1526221, 2025, doi: 10.3389/frai.2025.1526221.",
    "L. Breiman, “Random forests,” Machine Learning, vol. 45, pp. 5–32, 2001, doi: 10.1023/A:1010933404324.",
    "T. Chen and C. Guestrin, “XGBoost: A scalable tree boosting system,” in Proc. 22nd ACM SIGKDD Int. Conf. Knowledge Discovery and Data Mining, 2016, pp. 785–794, doi: 10.1145/2939672.2939785.",
]
for i, ref in enumerate(refs, 1):
    add_reference(doc, i, ref)

# Set all body sections to share the cleaned header and footer.
for section in doc.sections[1:]:
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True

# Ask Word to refresh fields.
settings = doc.settings._element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")

doc.save(OUTPUT)
print(OUTPUT)
